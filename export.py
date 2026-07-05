"""Excel export: a multi-sheet workbook reflecting the full historical
ledger, built with openpyxl. Read-only reporting — never writes back to the
DB. Deliberately depends only on db.py (not server.py) to avoid any import
cycle with the FastAPI route that serves it.
"""
import db
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _style_header(ws, row=1):
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _autosize(ws):
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(length + 2, 40)


def _write_table(ws, headers, rows):
    ws.append(headers)
    for row in rows:
        ws.append(row)
    _style_header(ws)
    ws.freeze_panes = "A2"
    _autosize(ws)


def _transactions_sheet(wb):
    ws = wb.active
    ws.title = "Transactions"
    txns = db.get_all_transactions()
    headers = ["Date", "Category", "Note", "Amount", "Payment Source", "Type", "Cadence", "Source", "Guessed"]
    rows = [
        [t["spent_on"], t["category"], t["note"] or "", t["amount"], t["payment_source"],
         t["expense_type"], t["cadence"] or "", t["source"], "yes" if t["guessed"] else ""]
        for t in txns
    ]
    _write_table(ws, headers, rows)


def _monthly_summary_sheet(wb):
    """One row per month with a full spend breakdown — the highest-level
    view of how the month-to-month picture has evolved over time."""
    ws = wb.create_sheet("Monthly Summary")
    months = db.get_all_months()
    headers = ["Month", "Total Spend", "Wallet Spend", "Credit Spend", "Liquid Spend",
               "Fixed", "Variable", "One-off", "Saving"]
    rows = []
    for month in months:
        txns = db.get_transactions_for_month(month)
        total = sum(t["amount"] for t in txns)
        wallet = sum(t["amount"] for t in txns if t["payment_source"] == "wallet")
        credit = sum(t["amount"] for t in txns if t["payment_source"] == "credit")
        liquid = sum(t["amount"] for t in txns if t["payment_source"] == "liquid")
        fixed = sum(t["amount"] for t in txns if t["expense_type"] == "fixed")
        variable = sum(t["amount"] for t in txns if t["expense_type"] == "variable")
        oneoff = sum(t["amount"] for t in txns if t["expense_type"] == "one-off")
        saving = sum(t["amount"] for t in txns if t["expense_type"] == "saving")
        rows.append([month, total, wallet, credit, liquid, fixed, variable, oneoff, saving])
    _write_table(ws, headers, rows)


def _category_breakdown_sheet(wb):
    """One row per month, one column per category — a wide pivot the user
    can filter/chart directly in Excel without reshaping the data first."""
    ws = wb.create_sheet("Category Breakdown")
    months = db.get_all_months()
    categories = sorted(c["name"] for c in db.get_categories())
    headers = ["Month"] + categories + ["Total"]
    rows = []
    for month in months:
        txns = db.get_transactions_for_month(month)
        by_cat = {}
        for t in txns:
            by_cat[t["category"]] = by_cat.get(t["category"], 0.0) + t["amount"]
        rows.append([month] + [by_cat.get(c, 0.0) for c in categories] + [sum(by_cat.values())])
    _write_table(ws, headers, rows)


def _accounts_ledger_sheet(wb):
    ws = wb.create_sheet("Accounts Ledger")
    txns = db.get_account_transactions(limit=1_000_000)
    headers = ["Date", "Account", "Type", "Amount", "Note"]
    rows = [
        [t["txn_date"], t["account"], t["txn_kind"], t["amount"], t["note"] or ""]
        for t in sorted(txns, key=lambda t: t["txn_date"])
    ]
    _write_table(ws, headers, rows)


def _investments_sheet(wb):
    """Long format: one row per (vehicle, month) with cumulative contributed
    and the manually-set value as of that month — the full history behind
    the dashboard's Savings section, for the user's own charting/analysis."""
    ws = wb.create_sheet("Investments")
    vehicles = [
        c for c in db.get_categories()
        if c["expense_type"] == "saving" and not c["account_link"]
    ]
    months = db.get_all_months()
    headers = ["Vehicle", "Month", "Contributed (cumulative)", "Value"]
    rows = []
    for cat in vehicles:
        name = cat["name"]
        contributed_by_month = {
            row["month"]: row["contributed"]
            for row in db.get_cumulative_savings_contributions(months, category=name)
        }
        for month in months:
            snap = db.get_snapshot_as_of(name, month)
            rows.append([name, month, contributed_by_month.get(month, 0.0), snap["value"] if snap else None])
    _write_table(ws, headers, rows)


def build_workbook():
    wb = Workbook()
    _transactions_sheet(wb)
    _monthly_summary_sheet(wb)
    _category_breakdown_sheet(wb)
    _accounts_ledger_sheet(wb)
    _investments_sheet(wb)
    return wb
